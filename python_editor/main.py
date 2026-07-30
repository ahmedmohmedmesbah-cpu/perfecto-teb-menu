import json
import shutil
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "data" / "products.json"
IMAGES_DIR = ROOT / "assets" / "images" / "products"
EXPORT_DIR = ROOT / "exports"
VALID_UNITS = ["قطعة", "كيلو"]
VALID_STATUSES = ["متوفر", "غير موجود حاليا"]
STATUS_FILTERS = ["All Products", "متوفر", "غير موجود حاليا"]
EXCEL_HEADERS = ["Product Name", "Product Price", "Selling Unit"]
EXCEL_FILENAME = "Perfecto_Products.xlsx"


class PerfectoEditor(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Perfecto Market - Product Editor")
        self.geometry("1040x680")
        self.minsize(920, 600)
        self.data = self.load_data()
        self.selected_product_id = None
        self.selected_category_id = None
        self.create_widgets()
        self.refresh_all()

    def load_data(self):
        if not DATA_FILE.exists():
            return {
                "meta": {
                    "marketNameAr": "بيرفكتو ماركت",
                    "marketNameEn": "Perfecto Market",
                    "currency": "جنيه",
                    "lastUpdated": datetime.now().strftime("%Y-%m-%d"),
                },
                "categories": [],
                "products": [],
            }

        with DATA_FILE.open("r", encoding="utf-8") as file:
            data = json.load(file)

        data.setdefault("meta", {})
        data.setdefault("categories", [])
        data.setdefault("products", [])

        for product in data["products"]:
            if "status" not in product:
                product["status"] = "متوفر" if product.get("available", True) else "غير موجود حاليا"
            if product["status"] not in VALID_STATUSES:
                product["status"] = "متوفر"

        return data

    def create_widgets(self):
        self.columnconfigure(0, weight=1)
        self.rowconfigure(1, weight=1)

        toolbar = ttk.Frame(self, padding=10)
        toolbar.grid(row=0, column=0, sticky="ew")
        toolbar.columnconfigure(7, weight=1)

        ttk.Button(toolbar, text="Import Excel", command=self.import_excel).grid(row=0, column=0, padx=4)
        ttk.Button(toolbar, text="Export All Excel", command=self.export_all_excel).grid(row=0, column=1, padx=4)
        ttk.Button(toolbar, text="Export Category Excel", command=self.export_selected_category_excel).grid(row=0, column=2, padx=4)
        ttk.Button(toolbar, text="Save JSON", command=self.save).grid(row=0, column=3, padx=4)
        ttk.Button(toolbar, text="Reload", command=self.reload).grid(row=0, column=4, padx=4)
        ttk.Button(toolbar, text="Add Category", command=self.new_category).grid(row=0, column=5, padx=4)
        ttk.Button(toolbar, text="Add Product", command=self.new_product).grid(row=0, column=6, padx=4)
        self.status_var = tk.StringVar(value=f"Editing: {DATA_FILE}")
        ttk.Label(toolbar, textvariable=self.status_var).grid(row=0, column=7, sticky="e")

        panes = ttk.PanedWindow(self, orient="horizontal")
        panes.grid(row=1, column=0, sticky="nsew")

        category_frame = ttk.Frame(panes, padding=10)
        category_frame.columnconfigure(0, weight=1)
        category_frame.rowconfigure(1, weight=1)
        panes.add(category_frame, weight=1)

        ttk.Label(category_frame, text="Categories").grid(row=0, column=0, sticky="w")
        self.category_tree = ttk.Treeview(category_frame, columns=("id", "name", "icon"), show="headings", height=12)
        for col, label, width in (("id", "ID", 140), ("name", "Arabic Name", 190), ("icon", "Icon", 90)):
            self.category_tree.heading(col, text=label)
            self.category_tree.column(col, width=width)
        self.category_tree.grid(row=1, column=0, sticky="nsew", pady=8)
        self.category_tree.bind("<<TreeviewSelect>>", self.on_category_select)

        category_form = ttk.LabelFrame(category_frame, text="Category Details", padding=10)
        category_form.grid(row=2, column=0, sticky="ew")
        category_form.columnconfigure(1, weight=1)

        self.cat_id = tk.StringVar()
        self.cat_name = tk.StringVar()
        self.cat_icon = tk.StringVar(value="basket")
        self.add_field(category_form, "ID", self.cat_id, 0)
        self.add_field(category_form, "Name", self.cat_name, 1)
        ttk.Label(category_form, text="Icon").grid(row=2, column=0, sticky="w", pady=4)
        ttk.Combobox(
            category_form,
            textvariable=self.cat_icon,
            values=("basket", "apple", "milk", "snow", "can", "sparkles"),
            state="readonly",
        ).grid(row=2, column=1, sticky="ew", pady=4)

        cat_actions = ttk.Frame(category_form)
        cat_actions.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(8, 0))
        ttk.Button(cat_actions, text="Apply Category", command=self.apply_category).pack(side="left", padx=3)
        ttk.Button(cat_actions, text="Delete Category", command=self.delete_category).pack(side="left", padx=3)

        product_frame = ttk.Frame(panes, padding=10)
        product_frame.columnconfigure(0, weight=1)
        product_frame.rowconfigure(3, weight=1)
        panes.add(product_frame, weight=2)

        filter_bar = ttk.Frame(product_frame)
        filter_bar.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        filter_bar.columnconfigure(1, weight=1)
        ttk.Label(filter_bar, text="Search").grid(row=0, column=0, padx=(0, 6))
        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", lambda *_: self.refresh_products())
        ttk.Entry(filter_bar, textvariable=self.search_var).grid(row=0, column=1, sticky="ew")
        ttk.Label(filter_bar, text="Status").grid(row=0, column=2, padx=(14, 6))
        self.status_filter = tk.StringVar(value=STATUS_FILTERS[0])
        ttk.Combobox(
            filter_bar,
            textvariable=self.status_filter,
            values=STATUS_FILTERS,
            state="readonly",
            width=16,
        ).grid(row=0, column=3, sticky="w")
        self.status_filter.trace_add("write", lambda *_: self.refresh_products())

        bulk_bar = ttk.Frame(product_frame)
        bulk_bar.grid(row=1, column=0, sticky="ew", pady=(0, 8))
        ttk.Button(bulk_bar, text="Bulk Delete", command=self.bulk_delete).pack(side="left", padx=3)
        ttk.Button(bulk_bar, text="Bulk Change Status", command=self.bulk_change_status).pack(side="left", padx=3)
        ttk.Button(bulk_bar, text="Bulk Change Category", command=self.bulk_change_category).pack(side="left", padx=3)
        ttk.Button(bulk_bar, text="Bulk Export", command=self.bulk_export_selected).pack(side="left", padx=3)
        ttk.Button(bulk_bar, text="Bulk Duplicate", command=self.bulk_duplicate).pack(side="left", padx=3)

        self.product_tree = ttk.Treeview(
            product_frame,
            columns=("id", "name", "category", "price", "unit", "status"),
            show="headings",
            selectmode="extended",
        )
        for col, label, width in (
            ("id", "ID", 120),
            ("name", "Product Name", 220),
            ("category", "Category", 130),
            ("price", "Price", 80),
            ("unit", "Unit", 90),
            ("status", "Status", 120),
        ):
            self.product_tree.heading(col, text=label)
            self.product_tree.column(col, width=width)
        self.product_tree.grid(row=2, column=0, sticky="nsew", pady=8)
        self.product_tree.bind("<<TreeviewSelect>>", self.on_product_select)
        self.product_tree.tag_configure("متوفر", foreground="#1f7b22")
        self.product_tree.tag_configure("غير موجود حاليا", foreground="#5f5f5f")

        product_form = ttk.LabelFrame(product_frame, text="Product Details", padding=10)
        product_form.grid(row=3, column=0, sticky="ew")
        product_form.columnconfigure(1, weight=1)

        self.prod_id = tk.StringVar()
        self.prod_name = tk.StringVar()
        self.prod_category = tk.StringVar()
        self.prod_price = tk.StringVar()
        self.prod_unit = tk.StringVar()
        self.prod_image = tk.StringVar()
        self.prod_status = tk.StringVar(value=VALID_STATUSES[0])

        self.add_field(product_form, "ID", self.prod_id, 0)
        self.add_field(product_form, "Name", self.prod_name, 1)
        ttk.Label(product_form, text="Category").grid(row=2, column=0, sticky="w", pady=4)
        self.category_combo = ttk.Combobox(product_form, textvariable=self.prod_category, state="readonly")
        self.category_combo.grid(row=2, column=1, sticky="ew", pady=4)
        self.add_field(product_form, "Price", self.prod_price, 3)
        self.add_field(product_form, "Unit", self.prod_unit, 4)
        ttk.Label(product_form, text="Status").grid(row=5, column=0, sticky="w", pady=4)
        ttk.Combobox(
            product_form,
            textvariable=self.prod_status,
            values=VALID_STATUSES,
            state="readonly",
        ).grid(row=5, column=1, sticky="ew", pady=4)
        self.add_field(product_form, "Image Path", self.prod_image, 6)
        ttk.Button(product_form, text="Choose Image", command=self.choose_image).grid(row=6, column=2, padx=6)

        product_actions = ttk.Frame(product_form)
        product_actions.grid(row=7, column=0, columnspan=3, sticky="ew", pady=(8, 0))
        ttk.Button(product_actions, text="Apply Product", command=self.apply_product).pack(side="left", padx=3)
        ttk.Button(product_actions, text="Delete Product", command=self.delete_product).pack(side="left", padx=3)

    def add_field(self, parent, label, variable, row):
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", pady=4)
        ttk.Entry(parent, textvariable=variable).grid(row=row, column=1, sticky="ew", pady=4)

    def category_names(self):
        return [f"{cat['name']} ({cat['id']})" for cat in self.data["categories"]]

    def selected_category_from_combo(self):
        value = self.prod_category.get()
        if "(" in value and value.endswith(")"):
            return value.rsplit("(", 1)[1][:-1]
        return value

    def refresh_all(self):
        self.refresh_categories()
        self.refresh_products()
        self.category_combo["values"] = self.category_names()

    def refresh_categories(self):
        self.category_tree.delete(*self.category_tree.get_children())
        for cat in self.data["categories"]:
            self.category_tree.insert("", "end", iid=cat["id"], values=(cat["id"], cat["name"], cat.get("icon", "basket")))

    def refresh_products(self):
        query = self.search_var.get().strip().lower() if hasattr(self, "search_var") else ""
        status_filter = self.status_filter.get() if hasattr(self, "status_filter") else STATUS_FILTERS[0]
        self.product_tree.delete(*self.product_tree.get_children())
        category_lookup = {cat["id"]: cat["name"] for cat in self.data["categories"]}

        for product in self.data["products"]:
            if query and query not in product["name"].lower() and query not in product["id"].lower():
                continue
            if status_filter != "All Products" and product.get("status", "متوفر") != status_filter:
                continue
            self.product_tree.insert(
                "",
                "end",
                iid=product["id"],
                values=(
                    product["id"],
                    product["name"],
                    category_lookup.get(product.get("categoryId"), product.get("categoryId", "")),
                    product.get("price", 0),
                    product.get("unit", ""),
                    product.get("status", "متوفر"),
                ),
                tags=(product.get("status", "متوفر"),),
            )

    def on_category_select(self, _event=None):
        selection = self.category_tree.selection()
        if not selection:
            return
        self.selected_category_id = selection[0]
        category = next(cat for cat in self.data["categories"] if cat["id"] == self.selected_category_id)
        self.cat_id.set(category["id"])
        self.cat_name.set(category["name"])
        self.cat_icon.set(category.get("icon", "basket"))

    def on_product_select(self, _event=None):
        selection = self.product_tree.selection()
        if not selection:
            return
        self.selected_product_id = selection[0]
        product = next(item for item in self.data["products"] if item["id"] == self.selected_product_id)
        self.prod_id.set(product["id"])
        self.prod_name.set(product["name"])
        self.prod_price.set(str(product.get("price", 0)))
        self.prod_unit.set(product.get("unit", ""))
        self.prod_image.set(product.get("image", ""))
        self.prod_status.set(product.get("status", "متوفر"))
        category = next((cat for cat in self.data["categories"] if cat["id"] == product.get("categoryId")), None)
        self.prod_category.set(f"{category['name']} ({category['id']})" if category else product.get("categoryId", ""))

    def new_category(self):
        self.selected_category_id = None
        self.cat_id.set("")
        self.cat_name.set("")
        self.cat_icon.set("basket")

    def new_product(self):
        self.selected_product_id = None
        self.prod_id.set("")
        self.prod_name.set("")
        self.prod_price.set("")
        self.prod_unit.set("")
        self.prod_image.set("")
        self.prod_status.set(VALID_STATUSES[0])
        self.prod_category.set(self.category_names()[0] if self.category_names() else "")

    def apply_category(self):
        category_id = self.cat_id.get().strip()
        name = self.cat_name.get().strip()
        if not category_id or not name:
            messagebox.showerror("Missing data", "Category ID and name are required.")
            return
        if self.selected_category_id and self.selected_category_id != category_id:
            for product in self.data["products"]:
                if product.get("categoryId") == self.selected_category_id:
                    product["categoryId"] = category_id
        existing = next((cat for cat in self.data["categories"] if cat["id"] == self.selected_category_id or cat["id"] == category_id), None)
        payload = {"id": category_id, "name": name, "icon": self.cat_icon.get()}
        if existing:
            existing.update(payload)
        else:
            self.data["categories"].append(payload)
        self.selected_category_id = category_id
        self.refresh_all()

    def delete_category(self):
        if not self.selected_category_id:
            return
        if any(product.get("categoryId") == self.selected_category_id for product in self.data["products"]):
            messagebox.showerror("Category in use", "Move or delete products in this category first.")
            return
        if messagebox.askyesno("Delete category", "Delete the selected category?"):
            self.data["categories"] = [cat for cat in self.data["categories"] if cat["id"] != self.selected_category_id]
            self.new_category()
            self.refresh_all()

    def apply_product(self):
        product_id = self.prod_id.get().strip()
        name = self.prod_name.get().strip()
        category_id = self.selected_category_from_combo()
        status = self.prod_status.get()
        try:
            price = float(self.prod_price.get())
        except ValueError:
            messagebox.showerror("Invalid price", "Price must be a number.")
            return
        if not product_id or not name or not category_id:
            messagebox.showerror("Missing data", "Product ID, name, category, and status are required.")
            return
        if status not in VALID_STATUSES:
            messagebox.showerror("Invalid status", "Please select a valid product status.")
            return
        payload = {
            "id": product_id,
            "categoryId": category_id,
            "name": name,
            "price": price,
            "unit": self.prod_unit.get().strip(),
            "image": self.prod_image.get().strip(),
            "status": status,
            "available": status == "متوفر",
        }
        existing = next((item for item in self.data["products"] if item["id"] == self.selected_product_id or item["id"] == product_id), None)
        if existing:
            existing.update(payload)
        else:
            self.data["products"].append(payload)
        self.selected_product_id = product_id
        self.refresh_products()

    def delete_product(self):
        if not self.selected_product_id:
            return
        if messagebox.askyesno("Delete product", "Delete the selected product?"):
            self.data["products"] = [item for item in self.data["products"] if item["id"] != self.selected_product_id]
            self.new_product()
            self.refresh_products()

    def choose_image(self):
        path = filedialog.askopenfilename(
            title="Choose product image",
            filetypes=(("Image files", "*.png *.jpg *.jpeg *.webp"), ("All files", "*.*")),
        )
        if not path:
            return
        IMAGES_DIR.mkdir(parents=True, exist_ok=True)
        source = Path(path)
        destination = IMAGES_DIR / source.name
        if source.resolve() != destination.resolve():
            shutil.copy2(source, destination)
        self.prod_image.set(destination.relative_to(ROOT).as_posix())

    def validate_data(self):
        errors = []
        category_ids = set()
        for cat in self.data["categories"]:
            if not cat.get("id"):
                errors.append("Category missing ID.")
            if not cat.get("name"):
                errors.append(f"Category {cat.get('id', '<unknown>')} is missing a name.")
            if cat.get("id") in category_ids:
                errors.append(f"Duplicate category ID: {cat['id']}")
            category_ids.add(cat.get("id"))

        product_ids = set()
        for product in self.data["products"]:
            if not product.get("id"):
                errors.append("Product missing ID.")
            if product["id"] in product_ids:
                errors.append(f"Duplicate product ID: {product['id']}")
            product_ids.add(product["id"])
            if not product.get("name"):
                errors.append(f"Product {product['id']} is missing a name.")
            if not isinstance(product.get("price", None), (int, float)):
                errors.append(f"Product {product['id']} has an invalid price.")
            if product.get("unit") not in VALID_UNITS:
                errors.append(f"Product {product['id']} has invalid unit: {product.get('unit')}")
            if product.get("status") not in VALID_STATUSES:
                errors.append(f"Product {product['id']} has invalid status: {product.get('status')}")
            if product.get("categoryId") not in category_ids:
                errors.append(f"Product {product['id']} references a missing category: {product.get('categoryId')}")

        return errors

    def backup_data_file(self):
        if DATA_FILE.exists():
            backup = DATA_FILE.with_name(f"products.backup.{datetime.now().strftime('%Y%m%d-%H%M%S')}.json")
            shutil.copy2(DATA_FILE, backup)
            return backup
        return None

    def save(self):
        errors = self.validate_data()
        if errors:
            messagebox.showerror("Validation failed", "\n".join(errors[:20]))
            return
        DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
        backup = self.backup_data_file()
        self.data.setdefault("meta", {})["lastUpdated"] = datetime.now().strftime("%Y-%m-%d")
        with DATA_FILE.open("w", encoding="utf-8") as file:
            json.dump(self.data, file, ensure_ascii=False, indent=2)
            file.write("\n")
        self.status_var.set(f"Saved: {DATA_FILE}")
        backup_text = f"\nBackup created: {backup.name}" if backup else ""
        messagebox.showinfo("Saved", f"products.json has been saved.{backup_text}")

    def export_selected_category_excel(self):
        if not self.selected_category_id:
            messagebox.showwarning("No category selected", "Select a category before export.")
            return
        category = next((cat for cat in self.data["categories"] if cat["id"] == self.selected_category_id), None)
        if not category:
            messagebox.showerror("Category missing", "The selected category could not be found.")
            return
        filename = filedialog.asksaveasfilename(
            title="Export Category Excel",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")],
            initialfile=f"Perfecto_{category['id']}.xlsx",
        )
        if not filename:
            return
        self.write_excel_workbook(filename, [category])
        messagebox.showinfo("Export complete", f"Excel exported to {filename}")

    def export_all_excel(self):
        if not self.data["categories"]:
            messagebox.showwarning("No categories", "Add categories before exporting.")
            return
        filename = filedialog.asksaveasfilename(
            title="Export All Categories",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")],
            initialfile=EXCEL_FILENAME,
        )
        if not filename:
            return
        self.write_excel_workbook(filename, self.data["categories"])
        messagebox.showinfo("Export complete", f"Excel exported to {filename}")

    def bulk_export_selected(self):
        selected_ids = self.product_tree.selection()
        if not selected_ids:
            messagebox.showwarning("No products selected", "Select products to export.")
            return
        products = [product for product in self.data["products"] if product["id"] in selected_ids]
        if not products:
            messagebox.showerror("Export error", "No matching products were found.")
            return
        filename = filedialog.asksaveasfilename(
            title="Export Selected Products",
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")],
            initialfile="Perfecto_Selected_Products.xlsx",
        )
        if not filename:
            return
        category_map = {cat["id"]: cat for cat in self.data["categories"]}
        categories = []
        for product in products:
            cat = category_map.get(product["categoryId"])
            if cat and cat not in categories:
                categories.append(cat)
        self.write_excel_workbook(filename, categories, products)
        messagebox.showinfo("Export complete", f"Excel exported to {filename}")

    def write_excel_workbook(self, filename, categories, products=None):
        workbook = Workbook()
        workbook.remove(workbook.active)

        for index, category in enumerate(categories, start=1):
            sheet_name = category["name"][:31]
            ws = workbook.create_sheet(title=sheet_name)
            ws.append(EXCEL_HEADERS)
            if products is None:
                category_products = [p for p in self.data["products"] if p.get("categoryId") == category["id"]]
            else:
                category_products = [p for p in products if p.get("categoryId") == category["id"]]
            for product in category_products:
                ws.append([product["name"], product.get("price", 0), product.get("unit", "")])
            table = Table(displayName=f"Table{index}", ref=f"A1:C{len(category_products) + 1}")
            style = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            table.tableStyleInfo = style
            ws.add_table(table)
            for col in range(1, 4):
                max_length = 0
                for row in ws.iter_rows(min_row=1, max_row=len(category_products) + 1, min_col=col, max_col=col, values_only=True):
                    cell_value = str(row[0] or "")
                    max_length = max(max_length, len(cell_value))
                ws.column_dimensions[get_column_letter(col)].width = max_length + 2
            for cell in ws[1]:
                cell.font = Font(bold=True)

        workbook.save(filename)

    def import_excel(self):
        filename = filedialog.askopenfilename(
            title="Import Excel Workbook",
            filetypes=[("Excel Workbook", "*.xlsx")],
        )
        if not filename:
            return
        self.backup_data_file()
        try:
            workbook = load_workbook(filename=filename, data_only=True)
        except Exception as ex:
            messagebox.showerror("Import failed", f"Unable to read Excel file: {ex}")
            return
        validation = self.validate_import_workbook(workbook)
        if validation["errors"]:
            self.show_import_summary(validation, allow_apply=False)
            return
        confirmed = self.show_import_summary(validation, allow_apply=True)
        if not confirmed:
            return
        self.apply_import_changes(validation)
        self.refresh_all()
        messagebox.showinfo(
            "Import complete",
            f"Import completed. {len(validation['products_updated'])} updated, {len(validation['products_added'])} added, {len(validation['products_skipped'])} skipped.",
        )

    def validate_import_workbook(self, workbook):
        summary = {
            "errors": [],
            "products_updated": [],
            "products_added": [],
            "products_skipped": [],
            "changes": [],
            "categories_created": [],
            "workbook_products": [],
        }
        existing_by_name = {product["name"]: product for product in self.data["products"]}
        existing_category_names = {cat["name"]: cat for cat in self.data["categories"]}
        seen_names = set()
        for sheet in workbook.worksheets:
            header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = [cell for cell in header_row[:3]]
            if headers != EXCEL_HEADERS:
                summary["errors"].append(f"Sheet '{sheet.title}' has invalid headers. Expected {EXCEL_HEADERS}.")
                continue
            category_name = sheet.title.strip()
            if not category_name:
                summary["errors"].append("Sheet with empty name is not allowed.")
                continue
            category = existing_category_names.get(category_name)
            if category is None:
                new_id = self.slugify(category_name)
                existing_ids = {cat["id"] for cat in self.data["categories"]}
                if new_id in existing_ids:
                    suffix = 2
                    while f"{new_id}-{suffix}" in existing_ids:
                        suffix += 1
                    new_id = f"{new_id}-{suffix}"
                category = {"id": new_id, "name": category_name, "icon": "basket"}
                existing_category_names[category_name] = category
                summary["categories_created"].append(category_name)
            for row_index, row in enumerate(sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=3, values_only=True), start=2):
                if row == (None, None, None):
                    continue
                name = str(row[0]).strip() if row[0] is not None else ""
                price_value = row[1]
                unit = str(row[2]).strip() if row[2] is not None else ""
                if not name:
                    summary["errors"].append(f"{sheet.title} row {row_index}: missing Product Name.")
                    continue
                if name in seen_names:
                    summary["errors"].append(f"Duplicate Product Name found in workbook: '{name}'.")
                    continue
                seen_names.add(name)
                try:
                    price = float(price_value)
                except (TypeError, ValueError):
                    summary["errors"].append(f"{sheet.title} row {row_index}: invalid price for '{name}'.")
                    continue
                if unit not in VALID_UNITS:
                    summary["errors"].append(f"{sheet.title} row {row_index}: invalid unit '{unit}' for '{name}'.")
                    continue
                summary["workbook_products"].append(
                    {
                        "name": name,
                        "price": price,
                        "unit": unit,
                        "category": category,
                        "sheet": sheet.title,
                    }
                )
        if summary["errors"]:
            return summary
        for row in summary["workbook_products"]:
            existing = existing_by_name.get(row["name"])
            if existing:
                changed = (
                    existing.get("price") != row["price"]
                    or existing.get("unit") != row["unit"]
                    or existing.get("categoryId") != row["category"]["id"]
                )
                if changed:
                    summary["products_updated"].append(row["name"])
                else:
                    summary["products_skipped"].append(row["name"])
            else:
                summary["products_added"].append(row["name"])
        return summary

    def apply_import_changes(self, summary):
        existing_by_name = {product["name"]: product for product in self.data["products"]}
        category_lookup = {cat["name"]: cat for cat in self.data["categories"]}
        for row in summary["workbook_products"]:
            category = category_lookup.get(row["category"]["name"])
            if category is None:
                category = row["category"]
                self.data["categories"].append(category)
                category_lookup[category["name"]] = category
            product = existing_by_name.get(row["name"])
            if product:
                product["price"] = row["price"]
                product["unit"] = row["unit"]
                product["categoryId"] = category["id"]
            else:
                new_id = self.create_unique_product_id(row["name"])
                new_product = {
                    "id": new_id,
                    "categoryId": category["id"],
                    "name": row["name"],
                    "price": row["price"],
                    "unit": row["unit"],
                    "image": "",
                    "status": "متوفر",
                    "available": True,
                }
                self.data["products"].append(new_product)

    def show_import_summary(self, summary, allow_apply):
        window = tk.Toplevel(self)
        window.title("Excel Import Summary")
        window.geometry("720x540")
        window.grab_set()

        text_frame = ttk.Frame(window, padding=12)
        text_frame.pack(fill="both", expand=True)

        summary_text = tk.Text(text_frame, wrap="word", state="normal")
        summary_text.pack(fill="both", expand=True)

        if summary["errors"]:
            summary_text.insert("end", "Errors detected during import:\n\n")
            for error in summary["errors"]:
                summary_text.insert("end", f"- {error}\n")
            summary_text.insert("end", "\nImport cannot continue until these errors are fixed.\n")
        else:
            summary_text.insert("end", f"Categories found: {len({row['category']['name'] for row in summary['workbook_products']})}\n")
            summary_text.insert("end", f"Products to update: {len(summary['products_updated'])}\n")
            summary_text.insert("end", f"Products to add: {len(summary['products_added'])}\n")
            summary_text.insert("end", f"Products skipped: {len(summary['products_skipped'])}\n")
            if summary["categories_created"]:
                summary_text.insert("end", "\nNew categories that will be created:\n")
                for category_name in summary["categories_created"]:
                    summary_text.insert("end", f"- {category_name}\n")

        summary_text.configure(state="disabled")

        button_frame = ttk.Frame(window, padding=12)
        button_frame.pack(fill="x")
        cancel_button = ttk.Button(button_frame, text="Cancel", command=window.destroy)
        cancel_button.pack(side="right", padx=6)

        if allow_apply and not summary["errors"]:
            result = {"confirmed": False}

            def apply_action():
                result["confirmed"] = True
                window.destroy()

            apply_button = ttk.Button(button_frame, text="Apply Changes", command=apply_action)
            apply_button.pack(side="right", padx=6)
            window.wait_window()
            return result["confirmed"]

        window.wait_window()
        return False

    def create_unique_product_id(self, name):
        base_id = self.slugify(name)
        product_ids = {product["id"] for product in self.data["products"]}
        candidate = base_id
        suffix = 1
        while candidate in product_ids:
            suffix += 1
            candidate = f"{base_id}-{suffix}"
        return candidate

    def slugify(self, text):
        value = str(text).strip().lower()
        value = value.replace(" ", "-")
        value = ''.join(ch for ch in value if ch.isalnum() or ch == "-")
        return value or "product"

    def bulk_delete(self):
        selected_ids = self.product_tree.selection()
        if not selected_ids:
            messagebox.showwarning("No products selected", "Select products to delete.")
            return
        if messagebox.askyesno("Delete selected products", f"Delete {len(selected_ids)} selected products?"):
            self.data["products"] = [p for p in self.data["products"] if p["id"] not in selected_ids]
            self.refresh_products()

    def bulk_change_status(self):
        selected_ids = self.product_tree.selection()
        if not selected_ids:
            messagebox.showwarning("No products selected", "Select products to update status.")
            return
        status = messagebox.askquestion("Product status", "Set status to 'متوفر' for selected products? Click No to set 'غير موجود حاليا'.")
        new_status = "متوفر" if status == "yes" else "غير موجود حاليا"
        for product in self.data["products"]:
            if product["id"] in selected_ids:
                product["status"] = new_status
                product["available"] = new_status == "متوفر"
        self.refresh_products()

    def bulk_change_category(self):
        selected_ids = self.product_tree.selection()
        if not selected_ids:
            messagebox.showwarning("No products selected", "Select products to change category.")
            return
        category_names = [cat["name"] for cat in self.data["categories"]]
        if not category_names:
            messagebox.showwarning("No categories", "Add categories first.")
            return
        category_window = tk.Toplevel(self)
        category_window.title("Bulk Change Category")
        category_window.geometry("360x140")
        category_window.grab_set()

        ttk.Label(category_window, text="Choose new category:").pack(anchor="w", padx=12, pady=(12, 4))
        selected_category = tk.StringVar(value=category_names[0])
        ttk.Combobox(category_window, textvariable=selected_category, values=category_names, state="readonly").pack(fill="x", padx=12)

        def apply_category():
            category_name = selected_category.get()
            category = next(cat for cat in self.data["categories"] if cat["name"] == category_name)
            for product in self.data["products"]:
                if product["id"] in selected_ids:
                    product["categoryId"] = category["id"]
            category_window.destroy()
            self.refresh_products()

        ttk.Button(category_window, text="Apply", command=apply_category).pack(side="right", padx=12, pady=12)
        ttk.Button(category_window, text="Cancel", command=category_window.destroy).pack(side="right", pady=12)
        category_window.wait_window()

    def bulk_duplicate(self):
        selected_ids = self.product_tree.selection()
        if not selected_ids:
            messagebox.showwarning("No products selected", "Select products to duplicate.")
            return
        for product in list(self.data["products"]):
            if product["id"] in selected_ids:
                copy_id = self.create_unique_product_id(f"{product['name']}-copy")
                duplicate = {
                    **product,
                    "id": copy_id,
                    "name": f"{product['name']} (نسخة)",
                }
                self.data["products"].append(duplicate)
        self.refresh_products()

    def reload(self):
        if messagebox.askyesno("Reload", "Discard unsaved changes and reload products.json?"):
            self.data = self.load_data()
            self.refresh_all()


if __name__ == "__main__":
    app = PerfectoEditor()
    app.mainloop()
